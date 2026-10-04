import hashlib
import shutil
import sqlite3
import sys
from datetime import date, datetime

from openpyxl import load_workbook

db_path = sys.argv[1] if len(sys.argv) > 1 else "cc.db"
xlsx_path = sys.argv[2] if len(sys.argv) > 2 else "Baseline_Carbon_Dioxide_Emission_Database_Version_22.0.xlsx"
backup = f"{db_path}.bak_{datetime.now():%Y%m%d_%H%M%S}"
shutil.copy2(db_path, backup)
print(f"Backup written: {backup}")

conn = sqlite3.connect(db_path)
conn.row_factory = sqlite3.Row

if not conn.execute("SELECT 1 FROM sqlite_master WHERE name = 'Data_Provenance'").fetchone():
    sys.exit("Data_Provenance not found - run migrate_carbonx.py first.")

# ---------------------------------------------------------------- read the CEA workbook
def read_cea(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    if "Results" not in wb.sheetnames:
        sys.exit(f"{path}: no 'Results' sheet - is this the CEA CO2 Baseline workbook?")
    rows = [[(j, v) for j, v in enumerate(r) if v is not None]
            for r in wb["Results"].iter_rows(values_only=True)]
    version = date_cell = None
    for r in rows:
        txt = [v for _, v in r]
        if txt and str(txt[0]).strip() == "VERSION":
            version = str(txt[1]).strip()
        if txt and str(txt[0]).strip() == "DATE":
            date_cell = txt[1]
    header = next((r for r in rows if r and str(r[0][1]).startswith("Emission Factors (tCO2/MWh) (excl. Imports)")), None)
    data = next((r for r in rows if r and str(r[0][1]).startswith("Weighted Average Grid Emission Rate (Incl. RES,Captive)")), None)
    if not (version and date_cell and header and data) or len(header) != 22 or len(data) != 22:
        sys.exit("Unexpected layout in the Results sheet - stopping so nothing wrong is imported.")
    years = [str(v) for _, v in header[12:22]]          # block INCLUDING imports
    values = [float(v) for _, v in data[12:22]]
    factors = {}
    for fy, val in zip(years, values):
        end_year = int(fy[:4]) + 1                         # 2024-25 -> 2025
        if not (0.3 < val < 1.5):
            sys.exit(f"Implausible factor {val} for {fy} - stopping.")
        factors[end_year] = val
    return version, date_cell.strftime("%B %Y"), factors


cea_ver, cea_month, CEA_FACTORS = read_cea(xlsx_path)
CEA_VERSION = f"{cea_ver} ({cea_month})"
sha256 = hashlib.sha256(open(xlsx_path, "rb").read()).hexdigest()
print(f"Read CEA workbook v{CEA_VERSION}, sha256 {sha256[:16]}...")
print("  FY-end year -> tCO2/MWh (incl. RES, captive, imports): " +
      ", ".join(f"{y}: {v:.4f}" for y, v in CEA_FACTORS.items()))

# ------------------------------------------------------------ 1
cols = [r["name"] for r in conn.execute("PRAGMA table_info(Emission_Record)")]
if "factor_id" not in cols:
    conn.execute("ALTER TABLE Emission_Record ADD COLUMN factor_id INTEGER "
                 "REFERENCES Emission_Factor(factor_id)")
# existing records were computed with whichever factor was active at the time
conn.execute("""
UPDATE Emission_Record
SET factor_id = (SELECT ef.factor_id
                 FROM Activity_Data ad
                 JOIN Emission_Factor ef ON ef.source_id = ad.source_id AND ef.is_active = 1
                 WHERE ad.activity_id = Emission_Record.activity_id)
WHERE factor_id IS NULL
""")

# ------------------------------------------------------------ 2 + 3
conn.executescript("""
DROP INDEX IF EXISTS idx_one_active_factor_per_source;
CREATE UNIQUE INDEX IF NOT EXISTS idx_one_active_factor_per_source_year
    ON Emission_Factor(source_id, effective_year) WHERE is_active = 1;

DROP TRIGGER IF EXISTS auto_emission_record;
CREATE TRIGGER auto_emission_record AFTER INSERT ON Activity_Data
FOR EACH ROW
BEGIN
    INSERT INTO Emission_Record (activity_id, emission_kg, factor_id)
    SELECT NEW.activity_id, NEW.quantity * ef.factor_value, ef.factor_id
    FROM Emission_Factor ef
    WHERE ef.factor_id = (
        SELECT f.factor_id FROM Emission_Factor f
        WHERE f.source_id = NEW.source_id AND f.is_active = 1
        ORDER BY (f.effective_year > CAST(NEW.period AS INTEGER)) ASC,
                 ABS(f.effective_year - CAST(NEW.period AS INTEGER)) ASC
        LIMIT 1);
END;
""")

# ------------------------------------------------------------ 4
src = conn.execute("SELECT source_id, unit FROM Emission_Source WHERE source_name = 'Electricity'").fetchone()
if src is None:
    sys.exit("No 'Electricity' row in Emission_Source.")
elec_id = src["source_id"]
if src["unit"] != "kWh":
    sys.exit(f"Electricity unit is {src['unit']!r}, expected 'kWh' (factor is kg CO2/kWh).")

METHODOLOGY = (
    "Read from the CEA workbook, sheet 'Results': 'Weighted Average Grid Emission Rate (Incl. RES,Captive)', "
    "block including cross-border imports, full precision. tCO2/MWh = kg CO2/kWh. "
    f"effective_year = fiscal-year end year (2025 = FY 2024-25). CO2 only. File SHA-256: {sha256}. "
    "Not used: the older 'Weighted Average Emission Rate' series, which excludes renewables and captive injection."
)
prov = conn.execute("SELECT provenance_id FROM Data_Provenance WHERE dataset_name = ? AND dataset_version = ?",
                    ("CEA CO2 Baseline Database for the Indian Power Sector", CEA_VERSION)).fetchone()
if prov is None:
    prov_id = conn.execute("""
        INSERT INTO Data_Provenance
            (dataset_name, organization, dataset_version, source_url, license,
             retrieval_date, methodology, is_synthetic)
        VALUES (?, ?, ?, ?, ?, ?, ?, 0)
    """, ("CEA CO2 Baseline Database for the Indian Power Sector",
          "Central Electricity Authority, Ministry of Power, Government of India",
          CEA_VERSION,
          "https://cea.nic.in/wp-content/uploads/baseline/2026/09/User_Guide__Version_22.0.pdf",
          "Not stated in the user guide - confirm terms of use with CEA before publication",
          date.today().isoformat(),
          METHODOLOGY)).lastrowid
else:
    prov_id = prov["provenance_id"]
    conn.execute("UPDATE Data_Provenance SET methodology = ? WHERE provenance_id = ?", (METHODOLOGY, prov_id))

retired = conn.execute("""
    UPDATE Emission_Factor SET is_active = 0
    WHERE source_id = ? AND is_active = 1
      AND provenance_id IN (SELECT provenance_id FROM Data_Provenance WHERE is_synthetic = 1)
""", (elec_id,)).rowcount
print(f"Retired {retired} unsourced electricity factor(s) (kept for history)")

for year, value in CEA_FACTORS.items():
    row = conn.execute("SELECT factor_id, factor_value FROM Emission_Factor WHERE source_id = ? "
                       "AND effective_year = ? AND provenance_id = ?", (elec_id, year, prov_id)).fetchone()
    if row is None:
        conn.execute("INSERT INTO Emission_Factor (source_id, factor_value, effective_year, is_active, provenance_id) "
                     "VALUES (?, ?, ?, 1, ?)", (elec_id, value, year, prov_id))
    elif abs(row["factor_value"] - value) > 1e-12:
        conn.execute("UPDATE Emission_Factor SET factor_value = ?, is_active = 1 WHERE factor_id = ?",
                     (value, row["factor_id"]))
print(f"Loaded CEA v{CEA_VERSION} grid factors for FY ending " + ", ".join(map(str, CEA_FACTORS)))

# ------------------------------------------------------------ 5
credit_before = {r["est_id"]: r["credit"] for r in conn.execute("SELECT est_id, credit FROM Carbon_Credit")}

print("\nElectricity emission records recomputed:")
acts = conn.execute("SELECT activity_id, est_id, period, quantity FROM Activity_Data WHERE source_id = ?",
                    (elec_id,)).fetchall()
for a in acts:
    f = conn.execute("""
        SELECT factor_id, factor_value, effective_year FROM Emission_Factor
        WHERE source_id = ? AND is_active = 1
        ORDER BY (effective_year > CAST(? AS INTEGER)) ASC,
                 ABS(effective_year - CAST(? AS INTEGER)) ASC
        LIMIT 1""", (elec_id, a["period"], a["period"])).fetchone()
    old = conn.execute("SELECT emission_kg FROM Emission_Record WHERE activity_id = ?",
                       (a["activity_id"],)).fetchone()
    new_kg = a["quantity"] * f["factor_value"]
    conn.execute("UPDATE Emission_Record SET emission_kg = ?, factor_id = ? WHERE activity_id = ?",
                 (new_kg, f["factor_id"], a["activity_id"]))
    print(f"  activity {a['activity_id']:>3}  period {a['period']}  {a['quantity']:>9.1f} kWh  "
          f"{(old['emission_kg'] if old else 0):>9.2f} -> {new_kg:>9.2f} kg  (factor {f['factor_value']:.4f}, FY end {f['effective_year']})")

conn.execute("""
UPDATE Carbon_Credit
SET credit = ROUND(COALESCE((SELECT SUM(credit_kg) FROM Credit_By_Period p
                             WHERE p.est_id = Carbon_Credit.est_id), 0) / 1000.0, 2)
""")
conn.execute("""
UPDATE Carbon_Credit
SET status = CASE WHEN credit > 0 THEN 'Surplus' WHEN credit < 0 THEN 'Deficit' ELSE 'Neutral' END
""")
conn.commit()

print("\nCredit changes (tCO2e):")
n = 0
for r in conn.execute("SELECT cc.est_id, e.est_name, cc.credit FROM Carbon_Credit cc JOIN Establishment e USING(est_id)"):
    if abs((credit_before[r["est_id"]] or 0) - (r["credit"] or 0)) > 1e-9:
        n += 1
        print(f"  {r['est_name']:<28} {credit_before[r['est_id']]:>8} -> {r['credit']:>8}")
if not n:
    print("  none")

# ------------------------------------------------------------ 6
print("\nVerification")
bad = conn.execute("""
    SELECT er.activity_id, er.emission_kg, ad.quantity * ef.factor_value AS expected
    FROM Emission_Record er
    JOIN Activity_Data ad ON ad.activity_id = er.activity_id
    LEFT JOIN Emission_Factor ef ON ef.factor_id = er.factor_id
    WHERE ef.factor_id IS NULL OR ABS(er.emission_kg - ad.quantity * ef.factor_value) > 0.01
""").fetchall()
print(f"  every emission record == quantity x its recorded factor: {'PASS' if not bad else 'FAIL'}")
for r in bad:
    print(f"    activity {r['activity_id']}: recorded {r['emission_kg']} vs expected {r['expected']}")

mism = conn.execute("""
    SELECT e.est_name, cc.credit, v.carbon_credit FROM Carbon_Credit cc
    JOIN Carbon_Report_View v ON v.est_id = cc.est_id JOIN Establishment e ON e.est_id = cc.est_id
    WHERE ABS(cc.credit - v.carbon_credit) > 0.011""").fetchall()
print(f"  stored credit == report view: {'PASS' if not mism else 'FAIL'}")

print("\nNOTE: Coal, Diesel, Natural Gas and Petrol factors are still unsourced seed values.")
conn.close()
print("\nDONE" if not bad and not mism else "\nDONE WITH FAILURES - restore from backup")